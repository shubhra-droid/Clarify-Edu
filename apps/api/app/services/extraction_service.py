"""Document extraction service — parses PDF, DOCX, and TXT files."""

import io
import logging
from pathlib import Path
import pdfplumber
import PyPDF2
import docx

logger = logging.getLogger(__name__)


class DocumentExtractionService:
    def extract_text(self, file_path: str | Path, mime_type: str) -> str:
        """Extract plain text from a file based on its MIME type."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found at {file_path}")

        logger.info("Extracting text from %s (MIME: %s)", path.name, mime_type)

        if mime_type == "text/plain" or path.suffix.lower() == ".txt":
            return self._extract_txt(path)
        elif mime_type == "application/pdf" or path.suffix.lower() == ".pdf":
            return self._extract_pdf(path)
        elif (
            mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            or path.suffix.lower() == ".docx"
        ):
            return self._extract_docx(path)
        else:
            raise ValueError(f"Unsupported file type for extraction: {mime_type}")

    def _extract_txt(self, path: Path) -> str:
        """Extract text from plain text file with encoding fallback."""
        try:
            with open(path, "r", encoding="utf-8") as f:
                return f.read()
        except UnicodeDecodeError:
            logger.warning("UTF-8 decoding failed for %s, falling back to ISO-8859-1", path.name)
            with open(path, "r", encoding="iso-8859-1") as f:
                return f.read()

    def _extract_pdf(self, path: Path) -> str:
        """Extract text from PDF using pdfplumber with PyPDF2 fallback."""
        text_content = []
        try:
            with pdfplumber.open(path) as pdf:
                for i, page in enumerate(pdf.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text_content.append(page_text)
                    else:
                        logger.debug("pdfplumber extracted empty text on page %d of %s", i + 1, path.name)
            
            full_text = "\n".join(text_content).strip()
            if full_text:
                return full_text
            logger.warning("pdfplumber returned empty text for %s. Falling back to PyPDF2.", path.name)
        except Exception as exc:
            logger.warning("pdfplumber extraction failed for %s: %s. Falling back to PyPDF2.", path.name, exc)

        # PyPDF2 fallback
        try:
            with open(path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                text_content = []
                for i, page in enumerate(reader.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text_content.append(page_text)
                full_text = "\n".join(text_content).strip()
                if full_text:
                    return full_text
        except Exception as exc:
            logger.error("PyPDF2 extraction fallback failed for %s: %s", path.name, exc)

        raise ValueError(f"Failed to extract readable text from PDF file: {path.name}")

    def _extract_docx(self, path: Path) -> str:
        """Extract text from Word Document using python-docx."""
        try:
            doc = docx.Document(path)
            paragraphs = [p.text for p in doc.paragraphs]
            
            # Extract text from tables too
            table_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text for cell in row.cells if cell.text]
                    if row_text:
                        table_text.append(" | ".join(row_text))
            
            full_content = []
            if paragraphs:
                full_content.append("\n".join(paragraphs))
            if table_text:
                full_content.append("\n=== Table Data ===\n" + "\n".join(table_text))
                
            return "\n\n".join(full_content).strip()
        except Exception as exc:
            logger.error("DOCX extraction failed for %s: %s", path.name, exc)
            raise ValueError(f"Failed to extract text from DOCX file: {path.name}") from exc


extraction_service = DocumentExtractionService()
